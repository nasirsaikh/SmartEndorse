import csv
import html
import io
import json
from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import httpx
from django.conf import settings
from django.core.mail import EmailMultiAlternatives
from django.db import transaction
from django.db.models import Q
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from pypdf import PdfReader

from .ai import AIService, CANONICAL_FIELDS
from .models import (
    Attachment, EndorsementApproval, EndorsementItem, EndorsementQuery,
    EndorsementRequest, IntegrationEndpoint, Organization, PlatformConfiguration,
    Policy, PolicyMember, PolicyPlan, PortalNotification, RecoveryUpload, UserProfile, WorkflowEvent,
)


def platform_config():
    obj, _ = PlatformConfiguration.objects.get_or_create(name="Default")
    return obj


def json_safe(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if hasattr(value, "pk"):
        return {"id": value.pk, "display": str(value)}
    if isinstance(value, dict):
        return {str(k): json_safe(v) for k, v in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [json_safe(v) for v in value]
    return value


class SLAService:
    @staticmethod
    def deadline(profile):
        if not profile:
            return None
        return timezone.now() + timedelta(hours=profile.target_hours)

    @classmethod
    def for_status(cls, request, status):
        if status in {EndorsementRequest.Status.SENT_TO_TPA, EndorsementRequest.Status.TPA_IN_PROGRESS}:
            return cls.deadline(request.policy.tpa_sla)
        if status in {EndorsementRequest.Status.TPA_QUERY, EndorsementRequest.Status.NEEDS_INFO}:
            return cls.deadline(request.policy.client_query_sla)
        if status in {
            EndorsementRequest.Status.SUBMITTED,
            EndorsementRequest.Status.PENDING_INSURER_APPROVAL,
            EndorsementRequest.Status.PENDING_AMOUNT_APPROVAL,
        }:
            return cls.deadline(request.policy.insurer_sla)
        return None


class NotificationService:
    STATUS_LEVELS = {
        EndorsementRequest.Status.COMPLETED: PortalNotification.Level.SUCCESS,
        EndorsementRequest.Status.NEEDS_INFO: PortalNotification.Level.WARNING,
        EndorsementRequest.Status.TPA_QUERY: PortalNotification.Level.WARNING,
        EndorsementRequest.Status.PENDING_INSURER_APPROVAL: PortalNotification.Level.WARNING,
        EndorsementRequest.Status.PENDING_AMOUNT_APPROVAL: PortalNotification.Level.WARNING,
        EndorsementRequest.Status.FAILED: PortalNotification.Level.DANGER,
        EndorsementRequest.Status.REJECTED: PortalNotification.Level.DANGER,
    }

    @staticmethod
    def _case_url(request_obj):
        return reverse("endorsement_detail", kwargs={"pk": request_obj.pk})

    @classmethod
    def recipients_for_request(cls, request_obj, include_tpa=True):
        users = {request_obj.requester}
        org_ids = {request_obj.policy.insurer_id, request_obj.requester_organization_id}
        if include_tpa and request_obj.policy.tpa_id:
            org_ids.add(request_obj.policy.tpa_id)
        for profile in UserProfile.objects.filter(organization_id__in=org_ids).select_related("user"):
            if profile.user.is_active:
                users.add(profile.user)
        return [u for u in users if u and u.is_active]

    @classmethod
    def create_portal(cls, users, title, message, request_obj=None, level=PortalNotification.Level.INFO, payload=None):
        cfg = platform_config()
        if not cfg.enable_portal_notifications:
            return
        link = cls._case_url(request_obj) if request_obj else ""
        PortalNotification.objects.bulk_create([
            PortalNotification(user=u, title=title, message=message, level=level, link=link, payload=payload or {})
            for u in set(users)
        ])

    @classmethod
    def _send_email(cls, users, subject, request_obj, heading, message):
        cfg = platform_config()
        if not cfg.enable_email_notifications:
            return
        emails = sorted({u.email for u in users if u.email})
        org_emails = {
            request_obj.policy.insurer.notification_email,
            request_obj.requester_organization.notification_email,
            request_obj.policy.tpa.notification_email if request_obj.policy.tpa_id else "",
        }
        emails.extend(sorted(e for e in org_emails if e and e not in emails))
        if not emails:
            return
        rows = "".join(
            f"<tr><td>{html.escape(i.full_name or i.member_no or str(i.pk))}</td>"
            f"<td>{html.escape(i.card_number or '—')}</td>"
            f"<td style='text-align:right'>{request_obj.currency} {i.premium_impact:.3f}</td></tr>"
            for i in request_obj.items.all()[:50]
        )
        body = f"""
        <div style="font-family:Arial,sans-serif;color:#172033;max-width:760px;margin:auto">
          <div style="padding:18px 22px;background:#15233d;color:#fff;border-radius:12px 12px 0 0">
            <div style="font-size:12px;opacity:.8">SMARTENDORSE</div><h2 style="margin:5px 0 0">{html.escape(heading)}</h2>
          </div>
          <div style="padding:22px;border:1px solid #e3e7ee;border-top:0;border-radius:0 0 12px 12px">
            <p>{html.escape(message)}</p>
            <table style="width:100%;border-collapse:collapse;margin:16px 0">
              <tr><td><b>Reference</b></td><td>{request_obj.reference}</td><td><b>Status</b></td><td>{html.escape(request_obj.get_status_display())}</td></tr>
              <tr><td><b>Policy</b></td><td>{html.escape(request_obj.policy.policy_number)}</td><td><b>Client</b></td><td>{html.escape(request_obj.policy.client.name)}</td></tr>
              <tr><td><b>Type</b></td><td>{html.escape(request_obj.get_endorsement_type_display())}</td><td><b>Effective</b></td><td>{request_obj.effective_date:%d %b %Y}</td></tr>
              <tr><td><b>Premium impact</b></td><td>{request_obj.currency} {request_obj.premium_impact:.3f}</td><td><b>SLA due</b></td><td>{request_obj.current_sla_due_at.strftime('%d %b %Y %H:%M') if request_obj.current_sla_due_at else '—'}</td></tr>
            </table>
            <h3 style="font-size:15px">Members</h3>
            <table style="width:100%;border-collapse:collapse"><thead><tr><th align="left">Member</th><th align="left">Card no.</th><th align="right">Impact</th></tr></thead><tbody>{rows or '<tr><td colspan="3">No rows</td></tr>'}</tbody></table>
            <p style="margin-top:18px;color:#5b6575">This email contains the information needed to understand the case. Portal access is only required when an action or correction is needed.</p>
          </div>
        </div>"""
        text = (
            f"{heading}\n\n{message}\nReference: {request_obj.reference}\nPolicy: {request_obj.policy.policy_number}\n"
            f"Status: {request_obj.get_status_display()}\nPremium impact: {request_obj.currency} {request_obj.premium_impact:.3f}"
        )
        email = EmailMultiAlternatives(subject=subject, body=text, from_email=settings.DEFAULT_FROM_EMAIL, to=emails)
        email.attach_alternative(body, "text/html")
        email.send(fail_silently=True)

    @classmethod
    def status_changed(cls, request_obj, description):
        if request_obj.status == EndorsementRequest.Status.VALIDATING:
            return
        users = cls.recipients_for_request(request_obj)
        title = f"{request_obj.reference} · {request_obj.get_status_display()}"
        level = cls.STATUS_LEVELS.get(request_obj.status, PortalNotification.Level.INFO)
        cls.create_portal(users, title, description or request_obj.get_status_display(), request_obj, level)
        cls._send_email(users, title, request_obj, request_obj.get_status_display(), description or "Endorsement status updated.")

    @classmethod
    def approval_requested(cls, approval):
        users = [p.user for p in UserProfile.objects.filter(organization=approval.assigned_organization).select_related("user") if p.user.is_active]
        users.append(approval.request.requester)
        title = f"Approval required · {approval.request.reference}"
        message = approval.reason or approval.get_approval_type_display()
        cls.create_portal(users, title, message, approval.request, PortalNotification.Level.WARNING, {"approval_id": approval.pk})
        cls._send_email(users, title, approval.request, "Approval required", message)


class PricingEngine:
    MONEY = Decimal("0.001")

    @classmethod
    def annual_premium(cls, policy, item):
        params = policy.rating_parameters or {}
        if policy.rating_method == Policy.RatingMethod.FLAT_ANNUAL:
            if item.plan_id:
                rate = item.plan.annual_rate
                rel_rates = item.plan.relationship_rates or {}
                if item.relationship and str(item.relationship) in rel_rates:
                    rate = Decimal(str(rel_rates[str(item.relationship)]))
                return Decimal(rate)
            return Decimal(str(params.get("default_annual_rate", "0")))
        if policy.rating_method == Policy.RatingMethod.PER_MILLE_SUM_ASSURED:
            return Decimal(item.sum_assured or 0) * Decimal(str(params.get("rate_per_mille", "0"))) / Decimal("1000")
        if policy.rating_method == Policy.RatingMethod.PERCENT_OF_SALARY:
            return Decimal(item.annual_salary or 0) * Decimal(str(params.get("salary_percent", "0"))) / Decimal("100")
        return Decimal("0")

    @classmethod
    def prorata_factor(cls, policy, effective_date):
        effective_date = max(effective_date, policy.effective_from)
        if effective_date > policy.effective_to:
            return Decimal("0")
        remaining = Decimal((policy.effective_to - effective_date).days + 1)
        mode = (policy.rating_parameters or {}).get("prorata_mode", "fixed_basis")
        denominator = Decimal((policy.effective_to - policy.effective_from).days + 1) if mode == "policy_days" else Decimal(policy.day_count_basis or 365)
        return min(Decimal("1"), max(Decimal("0"), remaining / denominator))

    @classmethod
    def calculate_item(cls, item):
        annual = cls.annual_premium(item.request.policy, item)
        factor = cls.prorata_factor(item.request.policy, item.effective_date or item.request.effective_date)
        impact = (annual * factor).quantize(cls.MONEY, rounding=ROUND_HALF_UP)
        if item.request.endorsement_type == EndorsementRequest.Type.DELETION:
            impact = -impact
        item.annual_premium = annual.quantize(cls.MONEY, rounding=ROUND_HALF_UP)
        item.prorata_factor = factor.quantize(Decimal("0.000001"), rounding=ROUND_HALF_UP)
        item.premium_impact = impact
        item.save(update_fields=["annual_premium", "prorata_factor", "premium_impact", "updated_at"])
        return impact

    @classmethod
    def calculate_request(cls, request_obj):
        total = sum((cls.calculate_item(item) for item in request_obj.items.select_related("plan")), Decimal("0"))
        request_obj.premium_impact = total.quantize(cls.MONEY, rounding=ROUND_HALF_UP)
        request_obj.currency = request_obj.policy.currency
        request_obj.save(update_fields=["premium_impact", "currency", "updated_at"])
        return request_obj.premium_impact


class FileIntakeService:
    HEADER_MAP = {
        "member no": "member_no", "member number": "member_no", "member id": "member_no", "membership no": "member_no",
        "employee no": "employee_no", "employee number": "employee_no", "employee id": "employee_no", "emp no": "employee_no",
        "national id": "national_id", "civil id": "national_id", "civil number": "national_id", "id no": "national_id",
        "name": "full_name", "full name": "full_name", "member name": "full_name", "insured name": "full_name",
        "relationship": "relationship", "relation": "relationship", "relation type": "relationship",
        "date of birth": "date_of_birth", "dob": "date_of_birth", "birth date": "date_of_birth",
        "gender": "gender", "sex": "gender",
        "plan": "plan_code", "plan code": "plan_code", "category": "plan_code", "class": "plan_code",
        "annual salary": "annual_salary", "salary": "annual_salary", "yearly salary": "annual_salary",
        "sum assured": "sum_assured", "sum insured": "sum_assured", "coverage amount": "sum_assured",
        "effective date": "effective_date", "addition date": "effective_date", "deletion date": "effective_date", "endorsement date": "effective_date",
    }

    @classmethod
    def kind_for_name(cls, name):
        ext = Path(name).suffix.lower()
        if ext in {".xlsx", ".xls", ".csv"}:
            return Attachment.Kind.EXCEL
        if ext == ".pdf":
            return Attachment.Kind.PDF
        if ext in {".png", ".jpg", ".jpeg", ".webp"}:
            return Attachment.Kind.IMAGE
        return Attachment.Kind.OTHER

    @classmethod
    def process(cls, attachment):
        try:
            raw_rows, normalized_rows, metadata = cls._extract(attachment)
            if attachment.is_supplemental:
                result = cls._recover_items(attachment.request, normalized_rows, attachment)
            else:
                result = {"created_items": cls._create_items(attachment.request, normalized_rows, attachment)}
            attachment.extracted_payload = json_safe({
                "source_file": attachment.original_name,
                "raw_rows": raw_rows,
                "normalized_rows": normalized_rows,
                **metadata,
                **result,
            })
            attachment.processed = True
            attachment.processing_error = ""
        except Exception as exc:
            attachment.processing_error = str(exc)
            attachment.processed = False
        attachment.save(update_fields=["extracted_payload", "processed", "processing_error", "updated_at"])
        return attachment.processed

    @classmethod
    def _extract(cls, attachment):
        ext = Path(attachment.original_name).suffix.lower()
        if ext == ".xlsx":
            raw = cls._xlsx_rows(attachment.file.path)
            normalized, ai_used = cls._normalize_structured(raw)
            return raw, normalized, {"method": "xlsx", "ai_header_mapping": ai_used}
        if ext == ".xls":
            raw = cls._xls_rows(attachment.file.path)
            normalized, ai_used = cls._normalize_structured(raw)
            return raw, normalized, {"method": "xls", "ai_header_mapping": ai_used}
        if ext == ".csv":
            raw = cls._csv_rows(attachment.file.path)
            normalized, ai_used = cls._normalize_structured(raw)
            return raw, normalized, {"method": "csv", "ai_header_mapping": ai_used}
        if ext == ".pdf":
            return cls._pdf_rows(attachment.file.path)
        if ext in {".png", ".jpg", ".jpeg", ".webp"}:
            rows = AIService().extract_image_rows(attachment.file.path)
            normalized = [cls._normalize_ai_row(r) for r in rows]
            return [r.get("_source_raw", r) for r in rows], normalized, {"method": "vision"}
        raise ValueError("Unsupported file type. Use XLSX, XLS, CSV, PDF, PNG, JPG, JPEG or WEBP.")

    @classmethod
    def _xlsx_rows(cls, path):
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        values = list(ws.iter_rows(values_only=True))
        if not values:
            return []
        headers = [str(x or "").strip() or f"column_{i+1}" for i, x in enumerate(values[0])]
        return [json_safe(dict(zip(headers, row))) for row in values[1:] if any(v not in (None, "") for v in row)]

    @classmethod
    def _xls_rows(cls, path):
        try:
            import xlrd
        except ImportError as exc:
            raise RuntimeError("Legacy XLS support requires xlrd.") from exc
        book = xlrd.open_workbook(path)
        sheet = book.sheet_by_index(0)
        if sheet.nrows == 0:
            return []
        headers = [str(sheet.cell_value(0, c) or "").strip() or f"column_{c+1}" for c in range(sheet.ncols)]
        rows = []
        for r in range(1, sheet.nrows):
            values = [sheet.cell_value(r, c) for c in range(sheet.ncols)]
            if any(v not in (None, "") for v in values):
                rows.append(dict(zip(headers, values)))
        return json_safe(rows)

    @classmethod
    def _csv_rows(cls, path):
        with open(path, "r", encoding="utf-8-sig", newline="") as handle:
            return [json_safe(row) for row in csv.DictReader(handle) if any(v not in (None, "") for v in row.values())]

    @classmethod
    def _pdf_rows(cls, path):
        reader = PdfReader(path)
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        if text.strip():
            rows = AIService().extract_text_rows(text)
            normalized = [cls._normalize_ai_row(r) for r in rows]
            return [r.get("_source_raw", r) for r in rows], normalized, {"method": "pdf_text_ai", "page_count": len(reader.pages)}
        ai = AIService()
        ai._require_provider(vision=True)
        try:
            import fitz
        except ImportError as exc:
            raise RuntimeError("Scanned PDF OCR requires PyMuPDF and a vision-capable AI provider.") from exc
        doc = fitz.open(path)
        all_rows = []
        for page in doc:
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            all_rows.extend(ai.extract_image_bytes(pix.tobytes("png"), "image/png"))
        normalized = [cls._normalize_ai_row(r) for r in all_rows]
        return [r.get("_source_raw", r) for r in all_rows], normalized, {"method": "pdf_vision_ocr", "page_count": len(doc)}

    @classmethod
    def _normalize_structured(cls, raw_rows):
        deterministic = []
        unknown_headers = set()
        for raw in raw_rows:
            mapped = {"_source_raw": json_safe(raw)}
            for key, value in raw.items():
                canonical = cls._map_header(key)
                if canonical in CANONICAL_FIELDS:
                    mapped[canonical] = value
                else:
                    unknown_headers.add(str(key))
            deterministic.append(mapped)
        if unknown_headers and AIService().available:
            try:
                return [cls._normalize_ai_row(r) for r in AIService().normalize_structured_rows(raw_rows)], True
            except Exception:
                pass
        return [cls._normalize_ai_row(r) for r in deterministic], False

    @classmethod
    def _normalize_ai_row(cls, row):
        source = row.get("_source_raw") if isinstance(row, dict) else None
        result = {"_source_raw": json_safe(source or row)}
        if isinstance(row, dict):
            for key in CANONICAL_FIELDS:
                result[key] = json_safe(row.get(key))
        return result

    @classmethod
    def _map_header(cls, value):
        key = " ".join(str(value or "").strip().lower().replace("_", " ").replace("-", " ").split())
        return cls.HEADER_MAP.get(key, key.replace(" ", "_"))

    @staticmethod
    def _date(value):
        if not value:
            return None
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        text = str(value).strip()
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%d.%m.%Y"):
            try:
                return datetime.strptime(text[:10], fmt).date()
            except ValueError:
                continue
        return None

    @staticmethod
    def _decimal(value):
        if value in (None, ""):
            return None
        try:
            return Decimal(str(value).replace(",", "").strip())
        except Exception:
            return None

    @classmethod
    def _to_item_values(cls, request_obj, row):
        plan = None
        if row.get("plan_code"):
            plan = PolicyPlan.objects.filter(policy=request_obj.policy, code__iexact=str(row["plan_code"]).strip(), is_active=True).first()
        member_no = str(row.get("member_no") or "").strip()
        existing = None
        if request_obj.endorsement_type == EndorsementRequest.Type.DELETION and member_no:
            existing = PolicyMember.objects.filter(policy=request_obj.policy, member_no=member_no, is_active=True).select_related("plan").first()
        return {
            "member_no": member_no,
            "employee_no": str(row.get("employee_no") or (existing.employee_no if existing else "") or "").strip(),
            "national_id": str(row.get("national_id") or (existing.national_id if existing else "") or "").strip(),
            "full_name": str(row.get("full_name") or (existing.full_name if existing else "") or "").strip(),
            "relationship": str(row.get("relationship") or (existing.relationship if existing else "") or "").strip(),
            "date_of_birth": cls._date(row.get("date_of_birth")) or (existing.date_of_birth if existing else None),
            "gender": str(row.get("gender") or (existing.gender if existing else "") or "").strip(),
            "plan": plan or (existing.plan if existing else None),
            "annual_salary": cls._decimal(row.get("annual_salary")) if row.get("annual_salary") not in (None, "") else (existing.annual_salary if existing else None),
            "sum_assured": cls._decimal(row.get("sum_assured")) if row.get("sum_assured") not in (None, "") else (existing.sum_assured if existing else None),
            "effective_date": cls._date(row.get("effective_date")) or request_obj.effective_date,
        }

    @classmethod
    def _create_items(cls, request_obj, rows, attachment):
        count = 0
        for row in rows:
            values = cls._to_item_values(request_obj, row)
            EndorsementItem.objects.create(
                request=request_obj,
                extracted_data=json_safe({"normalized": row, "source_raw": row.get("_source_raw", {}), "source_attachment_id": attachment.pk}),
                **values,
            )
            count += 1
        return count

    @staticmethod
    def _candidate_match(request_obj, row):
        qs = request_obj.items.all()
        for field in ("national_id", "member_no", "employee_no"):
            value = str(row.get(field) or "").strip()
            if value:
                matches = list(qs.filter(**{f"{field}__iexact": value})[:2])
                if len(matches) == 1:
                    return matches[0], field
        name = str(row.get("full_name") or "").strip()
        dob = FileIntakeService._date(row.get("date_of_birth"))
        if name and dob:
            matches = list(qs.filter(full_name__iexact=name, date_of_birth=dob)[:2])
            if len(matches) == 1:
                return matches[0], "name+dob"
        return None, ""

    @classmethod
    def _recover_items(cls, request_obj, rows, attachment):
        resolved, created, unmatched = 0, 0, 0
        editable = ["member_no", "employee_no", "national_id", "full_name", "relationship", "date_of_birth", "gender", "plan", "annual_salary", "sum_assured", "effective_date"]
        for row in rows:
            item, matched_by = cls._candidate_match(request_obj, row)
            values = cls._to_item_values(request_obj, row)
            if not item:
                if any(values.get(k) for k in ("member_no", "employee_no", "national_id", "full_name")):
                    item = EndorsementItem.objects.create(
                        request=request_obj,
                        extracted_data=json_safe({"normalized": row, "source_raw": row.get("_source_raw", {}), "source_attachment_id": attachment.pk, "supplemental": True}),
                        **values,
                    )
                    created += 1
                    WorkflowEvent.objects.create(request=request_obj, event_type="SUPPLEMENTAL_ROW_CREATED", description=f"Supplemental upload created item {item.pk}.", payload={"item_id": item.pk, "source": attachment.original_name})
                else:
                    unmatched += 1
                continue
            before = {f: json_safe(getattr(item, f + "_id") if f == "plan" else getattr(item, f)) for f in editable}
            changed = {}
            for field in editable:
                value = values.get(field)
                current = getattr(item, field)
                if value not in (None, "") and current in (None, ""):
                    setattr(item, field, value)
                    changed[field] = json_safe(value.pk if field == "plan" and value else value)
            if changed:
                item.resolution_data = {**(item.resolution_data or {}), "last_recovery": {"matched_by": matched_by, "source": attachment.original_name, "fields": changed}}
                item.extracted_data = {**(item.extracted_data or {}), "supplemental_sources": [*((item.extracted_data or {}).get("supplemental_sources", [])), json_safe(row.get("_source_raw", row))]}
                item.save()
                resolved += 1
                WorkflowEvent.objects.create(request=request_obj, event_type="ITEM_RECOVERED", description=f"Item {item.pk} supplemented from {attachment.original_name}.", payload={"item_id": item.pk, "matched_by": matched_by, "before": before, "filled": changed})
        return {"resolved_items": resolved, "created_items": created, "unmatched_rows": unmatched}


class BulkRecoveryService:
    """Resolve missing fields for one endorsement only, without guessing ambiguous matches."""

    MATCH_FIELDS = ("national_id", "member_no", "employee_no")
    EDITABLE_FIELDS = ("member_no", "employee_no", "national_id", "full_name", "relationship", "date_of_birth", "gender", "plan", "annual_salary", "sum_assured", "effective_date")

    @classmethod
    def process(cls, upload, endorsement):
        try:
            raw_rows, normalized_rows, metadata = FileIntakeService._extract(upload)
            items = endorsement.items.select_related("request", "request__policy", "plan")
            resolved = ambiguous = unmatched = 0
            outcomes = []

            for row_number, row in enumerate(normalized_rows, start=1):
                item, matched_by, is_ambiguous = cls._match(items, row)
                if is_ambiguous:
                    ambiguous += 1
                    outcomes.append({"row": row_number, "status": "ambiguous", "matched_by": matched_by})
                    continue
                if not item:
                    unmatched += 1
                    outcomes.append({"row": row_number, "status": "unmatched"})
                    continue

                values = FileIntakeService._to_item_values(endorsement, row)
                before, changed = {}, {}
                for field in cls.EDITABLE_FIELDS:
                    current = getattr(item, field)
                    value = values.get(field)
                    if value not in (None, "") and current in (None, ""):
                        before[field] = json_safe(current)
                        setattr(item, field, value)
                        changed[field] = json_safe(value.pk if field == "plan" and value else value)

                if not changed:
                    outcomes.append({
                        "row": row_number,
                        "status": "matched_no_missing_fields",
                        "item_id": item.pk,
                        "matched_by": matched_by,
                    })
                    continue

                item.resolution_data = {
                    **(item.resolution_data or {}),
                    "last_bulk_recovery": {
                        "source": upload.reference,
                        "matched_by": matched_by,
                        "fields": changed,
                    },
                }
                item.extracted_data = {
                    **(item.extracted_data or {}),
                    "bulk_recovery_sources": [
                        *((item.extracted_data or {}).get("bulk_recovery_sources", [])),
                        json_safe(row.get("_source_raw", row)),
                    ],
                }
                item.save()
                resolved += 1

                WorkflowEvent.objects.create(
                    request=endorsement,
                    actor=upload.uploaded_by,
                    event_type="BULK_ITEM_RECOVERED",
                    description=f"Item {item.pk} recovered from bulk upload {upload.reference}.",
                    payload=json_safe({
                        "item_id": item.pk,
                        "upload": upload.reference,
                        "matched_by": matched_by,
                        "before": before,
                        "filled": changed,
                    }),
                )
                outcomes.append({
                    "row": row_number,
                    "status": "resolved",
                    "item_id": item.pk,
                    "matched_by": matched_by,
                    "filled": changed,
                })

            upload.resolved_count = resolved
            upload.ambiguous_count = ambiguous
            upload.unmatched_count = unmatched
            upload.status = RecoveryUpload.Status.PROCESSED if not ambiguous and not unmatched else RecoveryUpload.Status.PARTIAL
            upload.processing_error = ""
            upload.extracted_payload = json_safe({
                "target_request_id": endorsement.pk,
                "target_reference": endorsement.reference,
                "raw_rows": raw_rows,
                "normalized_rows": normalized_rows,
                "outcomes": outcomes,
                **metadata,
            })
            upload.save(update_fields=[
                "resolved_count", "ambiguous_count", "unmatched_count", "status",
                "processing_error", "extracted_payload", "updated_at",
            ])
            return resolved
        except Exception as exc:
            upload.status = RecoveryUpload.Status.FAILED
            upload.processing_error = str(exc)
            upload.extracted_payload = json_safe({
                **(upload.extracted_payload or {}),
                "target_request_id": endorsement.pk,
                "target_reference": endorsement.reference,
            })
            upload.save(update_fields=["status", "processing_error", "extracted_payload", "updated_at"])
            return 0

    @classmethod
    def _match(cls, items, row):
        for field in cls.MATCH_FIELDS:
            value = str(row.get(field) or "").strip()
            if not value:
                continue
            matches = list(items.filter(**{f"{field}__iexact": value})[:3])
            if len(matches) == 1:
                return matches[0], field, False
            if len(matches) > 1:
                return None, field, True
        name = str(row.get("full_name") or "").strip()
        dob = FileIntakeService._date(row.get("date_of_birth"))
        if name and dob:
            matches = list(items.filter(full_name__iexact=name, date_of_birth=dob)[:3])
            if len(matches) == 1:
                return matches[0], "name+dob", False
            if len(matches) > 1:
                return None, "name+dob", True
        return None, "", False


class ValidationService:
    MEDICAL_ADD_DEFAULTS = ["full_name", "date_of_birth", "gender", "relationship", "plan"]
    LIFE_ADD_DEFAULTS = ["full_name", "date_of_birth", "gender", "sum_assured"]
    DELETE_DEFAULTS = ["member_no", "effective_date"]

    @classmethod
    def policy_gate_errors(cls, request_obj):
        policy = request_obj.policy
        today = timezone.localdate()
        cfg = platform_config()
        cutoff = policy.endorsement_expiry_cutoff_days
        if cutoff is None:
            cutoff = cfg.endorsement_expiry_cutoff_days
        errors = []
        if not policy.is_active:
            errors.append("Policy is inactive and cannot accept endorsements.")
        if policy.effective_to < today:
            errors.append(f"Policy expired on {policy.effective_to:%d %b %Y}; endorsements cannot be processed.")
        elif cutoff and (policy.effective_to - today).days <= cutoff:
            errors.append(f"Policy expires in {(policy.effective_to - today).days} days. Endorsements are blocked within {cutoff} days of expiry by backend configuration.")
        if request_obj.effective_date < policy.effective_from or request_obj.effective_date > policy.effective_to:
            errors.append(f"Effective date must be within policy period {policy.effective_from} to {policy.effective_to}.")
        if request_obj.effective_date < today - timedelta(days=policy.allow_backdated_days):
            errors.append(f"Effective date exceeds the allowed backdating limit of {policy.allow_backdated_days} day(s).")
        return errors

    @classmethod
    def validate(cls, request_obj):
        errors = cls.policy_gate_errors(request_obj)
        policy = request_obj.policy
        if policy.product == Policy.Product.GROUP_MEDICAL and not policy.tpa_id:
            errors.append("Group Medical policy has no TPA configured.")
        items = list(request_obj.items.select_related("plan"))
        if not items:
            errors.append("No endorsement member rows were provided or extracted. Nothing will be processed.")
        if request_obj.endorsement_type == EndorsementRequest.Type.ADDITION:
            required = policy.required_fields_addition or (
                cls.MEDICAL_ADD_DEFAULTS if policy.product == Policy.Product.GROUP_MEDICAL else cls.LIFE_ADD_DEFAULTS
            )
        else:
            required = policy.required_fields_deletion or cls.DELETE_DEFAULTS

        seen = {"member_no": {}, "employee_no": {}, "national_id": {}}
        for item in items:
            for field in seen:
                value = str(getattr(item, field) or "").strip().lower()
                if value:
                    seen[field].setdefault(value, []).append(item.pk)

        for item in items:
            item_errors = []
            for field in required:
                value = item.plan_id if field in {"plan", "plan_code"} else getattr(item, field, None)
                if value in (None, ""):
                    item_errors.append(f"{field.replace('_', ' ').title()} is required.")
            for field, values in seen.items():
                value = str(getattr(item, field) or "").strip().lower()
                if value and len(values.get(value, [])) > 1:
                    item_errors.append(f"Duplicate {field.replace('_', ' ')} in this request ({getattr(item, field)}).")

            existing = None
            if request_obj.endorsement_type == EndorsementRequest.Type.ADDITION:
                q = Q()
                if item.member_no:
                    q |= Q(member_no__iexact=item.member_no)
                if item.employee_no:
                    q |= Q(employee_no__iexact=item.employee_no)
                if item.national_id:
                    q |= Q(national_id__iexact=item.national_id)
                if q:
                    existing = PolicyMember.objects.filter(policy=policy, is_active=True).filter(q).first()
                item.is_existing_record = bool(existing)
                item.requires_insurer_approval = bool(existing)
            elif item.member_no:
                existing = PolicyMember.objects.filter(policy=policy, member_no=item.member_no, is_active=True).first()
                if not existing:
                    item_errors.append(f"Active member {item.member_no} was not found on the policy.")

            item.validation_errors = item_errors
            if item_errors:
                item.validation_status = EndorsementItem.ValidationStatus.ERROR
            elif item.is_existing_record:
                item.validation_status = EndorsementItem.ValidationStatus.EXISTING
            else:
                item.validation_status = EndorsementItem.ValidationStatus.VALID
            item.save(update_fields=["validation_errors", "validation_status", "is_existing_record", "requires_insurer_approval", "updated_at"])
            errors.extend([f"Item {item.pk}: {e}" for e in item_errors])

        for attachment in request_obj.attachments.exclude(processing_error=""):
            errors.append(f"{attachment.original_name}: {attachment.processing_error}")

        mandatory = policy.mandatory_documents_addition if request_obj.endorsement_type == EndorsementRequest.Type.ADDITION else policy.mandatory_documents_deletion
        if mandatory:
            names = " ".join(request_obj.attachments.values_list("original_name", flat=True)).lower()
            for doc in mandatory:
                tokens = str(doc).lower().replace("_", " ").split()
                if tokens and not all(token in names for token in tokens):
                    errors.append(f"Mandatory document '{doc}' was not identified in uploaded filenames.")

        score = Decimal("100.00") if not errors else max(Decimal("0"), Decimal("100") - Decimal(len(errors) * 5))
        return errors, score


class DispatchService:
    @classmethod
    def dispatch(cls, request_obj):
        if request_obj.policy.product == Policy.Product.GROUP_MEDICAL:
            org, owner_type = request_obj.policy.tpa, IntegrationEndpoint.OwnerType.TPA
        else:
            org, owner_type = request_obj.policy.insurer, IntegrationEndpoint.OwnerType.INSURER_CORE
        endpoint = IntegrationEndpoint.objects.filter(organization=org, owner_type=owner_type, is_active=True).filter(Q(product="") | Q(product=request_obj.policy.product)).first()
        if not endpoint:
            raise RuntimeError(f"No active {owner_type} integration endpoint configured for {org}.")
        payload = {
            "reference": request_obj.reference,
            "policy_number": request_obj.policy.policy_number,
            "endorsement_type": request_obj.endorsement_type,
            "effective_date": request_obj.effective_date.isoformat(),
            "premium_impact": str(request_obj.premium_impact),
            "currency": request_obj.currency,
            "items": [{
                "member_no": i.member_no, "employee_no": i.employee_no, "national_id": i.national_id,
                "full_name": i.full_name, "relationship": i.relationship,
                "date_of_birth": i.date_of_birth.isoformat() if i.date_of_birth else None,
                "gender": i.gender, "plan": i.plan.code if i.plan else None,
                "annual_salary": str(i.annual_salary) if i.annual_salary is not None else None,
                "sum_assured": str(i.sum_assured) if i.sum_assured is not None else None,
                "premium_impact": str(i.premium_impact), "raw_extraction": i.extracted_data,
            } for i in request_obj.items.select_related("plan")],
        }
        if endpoint.transport == IntegrationEndpoint.Transport.API:
            response = httpx.post(endpoint.endpoint_url, json=payload, headers=endpoint.auth_headers or {}, timeout=60)
            response.raise_for_status()
            try:
                data = response.json()
                request_obj.external_reference = str(data.get("reference") or data.get("id") or "")
                request_obj.save(update_fields=["external_reference", "updated_at"])
            except Exception:
                pass
        else:
            recipient = endpoint.recipient_email or org.notification_email
            if not recipient:
                raise RuntimeError("Email integration endpoint has no recipient.")
            body = (
                f"New endorsement {request_obj.reference}\nPolicy: {request_obj.policy.policy_number}\n"
                f"Type: {request_obj.get_endorsement_type_display()}\nEffective: {request_obj.effective_date}\n"
                f"Members: {request_obj.items.count()}\nPremium impact: {request_obj.currency} {request_obj.premium_impact:.3f}\n"
                "All extracted source JSON is retained in SmartEndorse for audit."
            )
            EmailMultiAlternatives(
                subject=f"SmartEndorse {request_obj.reference} - {request_obj.get_endorsement_type_display()}",
                body=body, from_email=settings.DEFAULT_FROM_EMAIL, to=[recipient],
            ).send(fail_silently=False)
        return endpoint


class WorkflowService:
    @staticmethod
    def transition(request_obj, status, actor=None, description="", payload=None, notify=True):
        old = request_obj.status
        request_obj.status = status
        request_obj.current_sla_due_at = SLAService.for_status(request_obj, status)
        if status == EndorsementRequest.Status.COMPLETED:
            request_obj.completed_at = timezone.now()
            request_obj.current_sla_due_at = None
        request_obj.save(update_fields=["status", "current_sla_due_at", "completed_at", "updated_at"])
        WorkflowEvent.objects.create(
            request=request_obj, actor=actor, event_type="STATUS_CHANGE", from_status=old, to_status=status,
            description=description, payload=json_safe(payload or {}),
        )
        if notify:
            NotificationService.status_changed(request_obj, description)

    @classmethod
    @transaction.atomic
    def submit(cls, request_obj, actor=None):
        cls.transition(request_obj, EndorsementRequest.Status.VALIDATING, actor, "Validation started.", notify=False)
        errors, score = ValidationService.validate(request_obj)
        request_obj.validation_errors = errors
        request_obj.validation_score = score
        request_obj.submitted_at = timezone.now()
        request_obj.save(update_fields=["validation_errors", "validation_score", "submitted_at", "updated_at"])
        if errors or not request_obj.items.exists():
            cls.transition(request_obj, EndorsementRequest.Status.NEEDS_INFO, actor, "Validation found missing/invalid data. Processing is blocked until all errors are resolved.", {"errors": errors})
            return request_obj

        PricingEngine.calculate_request(request_obj)
        cfg = platform_config()
        approval_items = request_obj.items.filter(requires_insurer_approval=True)
        if approval_items.exists():
            request_obj.stp_eligible = False
            request_obj.save(update_fields=["stp_eligible", "updated_at"])
            for item in approval_items:
                approval, created = EndorsementApproval.objects.get_or_create(
                    request=request_obj, item=item, approval_type=EndorsementApproval.ApprovalType.EXISTING_MEMBER,
                    status=EndorsementApproval.Status.PENDING,
                    defaults={"assigned_organization": request_obj.policy.insurer, "requested_by": actor, "reason": f"Existing active member match detected for {item.full_name or item.member_no}. Insurer approval is required before processing."},
                )
                if created:
                    item.validation_status = EndorsementItem.ValidationStatus.APPROVAL_REQUIRED
                    item.save(update_fields=["validation_status", "updated_at"])
                    NotificationService.approval_requested(approval)
            cls.transition(request_obj, EndorsementRequest.Status.PENDING_INSURER_APPROVAL, actor, "Existing member record(s) require insurance company approval before processing.")
            return request_obj

        request_obj.stp_eligible = bool(cfg.auto_stp_enabled and request_obj.policy.auto_stp)
        request_obj.save(update_fields=["stp_eligible", "updated_at"])
        cls.transition(request_obj, EndorsementRequest.Status.SUBMITTED, actor, "Request passed validation and pricing.")
        if not request_obj.stp_eligible:
            return request_obj
        return cls._auto_dispatch(request_obj)

    @classmethod
    def _auto_dispatch(cls, request_obj):
        cls.transition(request_obj, EndorsementRequest.Status.AUTO_APPROVED, None, "Straight-through rules passed. No insurer manual approval required.")
        try:
            endpoint = DispatchService.dispatch(request_obj)
            target = EndorsementRequest.Status.SENT_TO_TPA if request_obj.policy.product == Policy.Product.GROUP_MEDICAL else EndorsementRequest.Status.CORE_DISPATCHED
            cls.transition(request_obj, target, None, f"Automatically dispatched through {endpoint.get_transport_display()}.")
        except Exception as exc:
            request_obj.validation_errors = [*request_obj.validation_errors, f"Automatic dispatch failed: {exc}"]
            request_obj.save(update_fields=["validation_errors", "updated_at"])
            cls.transition(request_obj, EndorsementRequest.Status.FAILED, None, "Automatic dispatch failed.", {"error": str(exc)})
        return request_obj

    @classmethod
    def revalidate_after_correction(cls, request_obj, actor):
        return cls.submit(request_obj, actor)

    @classmethod
    def start_processing(cls, request_obj, actor):
        gate = ValidationService.policy_gate_errors(request_obj)
        if gate:
            raise ValueError(" ".join(gate))
        if request_obj.validation_errors:
            raise ValueError("Validation errors must be resolved before processing.")
        cls.transition(request_obj, EndorsementRequest.Status.TPA_IN_PROGRESS, actor, "TPA processing started.")

    @classmethod
    def update_tpa_item(cls, item, actor, card_number, amount):
        request_obj = item.request
        if request_obj.policy.product != Policy.Product.GROUP_MEDICAL:
            raise ValueError("TPA card/amount processing applies only to Group Medical endorsements.")
        old_card, old_amount = item.card_number, item.tpa_premium_amount
        amount = Decimal(str(amount)) if amount not in (None, "") else item.premium_impact
        item.card_number = str(card_number or "").strip()
        item.tpa_premium_amount = amount
        item.save(update_fields=["card_number", "tpa_premium_amount", "updated_at"])
        WorkflowEvent.objects.create(
            request=request_obj, actor=actor, event_type="TPA_ITEM_UPDATED",
            description=f"TPA updated card/amount for item {item.pk}.",
            payload=json_safe({"item_id": item.pk, "before": {"card_number": old_card, "amount": old_amount}, "after": {"card_number": item.card_number, "amount": amount}}),
        )
        if amount != item.premium_impact:
            cfg = platform_config()
            assigned_org = request_obj.policy.insurer if cfg.tpa_amount_approval_party == PlatformConfiguration.AmountApprovalParty.INSURER else request_obj.policy.client
            approval = EndorsementApproval.objects.filter(
                request=request_obj, item=item, approval_type=EndorsementApproval.ApprovalType.TPA_AMOUNT_CHANGE,
                status=EndorsementApproval.Status.PENDING,
            ).first()
            if approval:
                approval.assigned_organization = assigned_org
                approval.requested_by = actor
                approval.reason = f"TPA changed amount from {request_obj.currency} {item.premium_impact:.3f} to {request_obj.currency} {amount:.3f} for {item.full_name or item.member_no}."
                approval.old_amount = item.premium_impact
                approval.new_amount = amount
                approval.save(update_fields=["assigned_organization", "requested_by", "reason", "old_amount", "new_amount", "updated_at"])
            else:
                approval = EndorsementApproval.objects.create(
                    request=request_obj, item=item, approval_type=EndorsementApproval.ApprovalType.TPA_AMOUNT_CHANGE,
                    assigned_organization=assigned_org, requested_by=actor,
                    reason=f"TPA changed amount from {request_obj.currency} {item.premium_impact:.3f} to {request_obj.currency} {amount:.3f} for {item.full_name or item.member_no}.",
                    old_amount=item.premium_impact, new_amount=amount,
                )
            cls.transition(request_obj, EndorsementRequest.Status.PENDING_AMOUNT_APPROVAL, actor, "TPA changed the calculated amount. Approval is required before completion.")
            NotificationService.approval_requested(approval)
            return approval
        return None

    @classmethod
    def decide_approval(cls, approval, actor, approve=True, comment=""):
        approval.status = EndorsementApproval.Status.APPROVED if approve else EndorsementApproval.Status.REJECTED
        approval.decided_by = actor
        approval.decided_at = timezone.now()
        approval.decision_comment = comment
        approval.save(update_fields=["status", "decided_by", "decided_at", "decision_comment", "updated_at"])
        request_obj = approval.request
        WorkflowEvent.objects.create(
            request=request_obj, actor=actor, event_type="APPROVAL_DECISION",
            description=f"{approval.get_approval_type_display()} {approval.get_status_display().lower()}.",
            payload=json_safe({"approval_id": approval.pk, "type": approval.approval_type, "decision": approval.status, "comment": comment}),
        )
        if not approve:
            if approval.approval_type == EndorsementApproval.ApprovalType.TPA_AMOUNT_CHANGE and approval.item_id:
                approval.item.tpa_premium_amount = approval.old_amount
                approval.item.save(update_fields=["tpa_premium_amount", "updated_at"])
                cls.transition(request_obj, EndorsementRequest.Status.TPA_IN_PROGRESS, actor, "TPA amount change rejected; system-calculated amount restored.")
            else:
                cls.transition(request_obj, EndorsementRequest.Status.REJECTED, actor, "Existing member exception was rejected by the insurer.")
            return request_obj

        if request_obj.approvals.filter(status=EndorsementApproval.Status.PENDING).exists():
            return request_obj
        if approval.approval_type == EndorsementApproval.ApprovalType.EXISTING_MEMBER:
            for item in request_obj.items.filter(requires_insurer_approval=True):
                item.validation_status = EndorsementItem.ValidationStatus.VALID
                item.save(update_fields=["validation_status", "updated_at"])
            cls.transition(request_obj, EndorsementRequest.Status.SUBMITTED, actor, "Existing-member exception approved by insurer; request released for downstream processing.")
            try:
                endpoint = DispatchService.dispatch(request_obj)
                target = EndorsementRequest.Status.SENT_TO_TPA if request_obj.policy.product == Policy.Product.GROUP_MEDICAL else EndorsementRequest.Status.CORE_DISPATCHED
                cls.transition(request_obj, target, actor, f"Approved exception dispatched through {endpoint.get_transport_display()}.")
            except Exception as exc:
                request_obj.validation_errors = [*request_obj.validation_errors, f"Dispatch after approval failed: {exc}"]
                request_obj.save(update_fields=["validation_errors", "updated_at"])
                cls.transition(request_obj, EndorsementRequest.Status.FAILED, actor, "Dispatch after approval failed.", {"error": str(exc)})
            return request_obj
        cls.transition(request_obj, EndorsementRequest.Status.TPA_IN_PROGRESS, actor, "Changed TPA amount approved. TPA may continue processing.")
        return request_obj

    @classmethod
    def complete(cls, request_obj, actor, external_reference=""):
        gate = ValidationService.policy_gate_errors(request_obj)
        if gate:
            raise ValueError(" ".join(gate))
        if request_obj.validation_errors or request_obj.items.filter(validation_status=EndorsementItem.ValidationStatus.ERROR).exists():
            raise ValueError("The endorsement still contains validation errors.")
        if request_obj.approvals.filter(status=EndorsementApproval.Status.PENDING).exists():
            raise ValueError("Pending approval(s) must be completed first.")
        if request_obj.policy.product == Policy.Product.GROUP_MEDICAL and request_obj.items.filter(card_number="").exists():
            raise ValueError("Card number is mandatory for every Medical endorsement item before completion.")
        if external_reference:
            request_obj.external_reference = external_reference
            request_obj.save(update_fields=["external_reference", "updated_at"])
        cls.transition(request_obj, EndorsementRequest.Status.COMPLETED, actor, "Endorsement completed. All required parties have been notified.")

    @classmethod
    def raise_query(cls, request_obj, actor, subject, message):
        request_obj.metadata = {**(request_obj.metadata or {}), "pre_query_status": request_obj.status}
        request_obj.save(update_fields=["metadata", "updated_at"])
        query = EndorsementQuery.objects.create(
            request=request_obj, raised_by=actor, assigned_organization=request_obj.requester_organization,
            subject=subject, message=message, due_at=SLAService.deadline(request_obj.policy.client_query_sla),
        )
        cls.transition(request_obj, EndorsementRequest.Status.TPA_QUERY, actor, "Additional information requested from requester.")
        return query

    @classmethod
    def answer_query(cls, query, actor, response):
        query.response = response
        query.responded_by = actor
        query.responded_at = timezone.now()
        query.is_closed = True
        query.save(update_fields=["response", "responded_by", "responded_at", "is_closed", "updated_at"])
        restore = (query.request.metadata or {}).get("pre_query_status") or EndorsementRequest.Status.TPA_IN_PROGRESS
        cls.transition(query.request, restore, actor, "Requester answered query.")
